"""Builds report/report.docx and report/slides.pptx from the results in report/ (csv + figures)."""
from pathlib import Path

import pandas as pd
from docx import Document
from docx.shared import Inches
from pptx import Presentation
from pptx.util import Inches as PI, Pt

R = Path(__file__).resolve().parent
F = R / "figures"
prof = pd.read_csv(R / "cluster_profile.csv")
rules = pd.read_csv(R / "churn_rules.csv").head(8)
imp = pd.read_csv(R / "feature_importance.csv", index_col=0).iloc[:, 0].head(6)

NAMES = {0: "Đã lỡ hạn, vẫn nghe", 1: "Thuê bao đang hoạt động (nhóm lớn)",
         2: "Trả phí nhưng không nghe", 3: "Hủy nhiều, đơn giá lạ", 4: "Đã bỏ hoàn toàn"}

SECTIONS = [
 ("1. Giới thiệu và Business Understanding", [
  "KKBox là dịch vụ nghe nhạc theo thuê bao. Doanh thu phụ thuộc vào việc người dùng gia hạn đúng hạn, nên hiểu ai rời đi và vì sao có giá trị hơn việc chỉ dự đoán.",
  "Bài toán: khám phá các nhóm hành vi và tổ hợp đặc điểm gắn với việc không gia hạn (churn) để thiết kế chiến lược giữ chân theo từng nhóm.",
  "Nhiệm vụ data mining: (1) Clustering để phân khúc người dùng; (2) Association rules (FP-Growth) để tìm tổ hợp dẫn tới churn; (3) Classification (Decision Tree, Random Forest) làm phương pháp bổ trợ để giải thích yếu tố quan trọng.",
  "Giả định: nhãn churn theo định nghĩa của competition (không gia hạn trong 30 ngày sau khi hết hạn). Hạn chế: dữ liệu là một lát cắt thời gian, mẫu được lấy mất cân bằng có chủ đích (xem mục 3), tuổi thiếu nhiều."]),
 ("2. Dataset", [
  "Nguồn: WSDM – KKBox's Churn Prediction Challenge (Kaggle), dữ liệu thật. Dùng 4 bảng: train_v2 (970.960 người dùng, 9,0% churn), members_v3 (thông tin thành viên), transactions_v2 (giao dịch đăng ký), user_logs_v2 (log nghe nhạc, ~1,4 GB).",
  "Không dùng user_logs.csv gốc (~28 GB) vì giới hạn ổ đĩa. Các bảng nối với nhau qua msno (ID người dùng đã hash).",
  "Thuộc tính: định danh/phân loại (city, gender, registered_via, payment_method), số (tuổi, giá, số ngày gói), thời gian (ngày đăng ký, ngày giao dịch), và số lượt nghe theo mức hoàn thành bài."]),
 ("3. Data Pre-processing", [
  "Cột bd (tuổi) có nhiều giá trị 0, âm hoặc > 100: giữ giá trị trong 10–80, còn lại coi là thiếu và gán bằng trung vị. Hệ quả: 55% người dùng mang tuổi được gán, nên tuổi gần như không dùng được để diễn giải.",
  "gender thiếu: gán 'unknown'. total_secs có giá trị phi thực tế (> 24 giờ/ngày): chặn ở 86.400 giây.",
  "Aggregate theo người dùng: số giao dịch, số lần hủy, tỉ lệ tự gia hạn, giá trung bình, số ngày hoạt động, tổng số bài, tỉ lệ nghe hết (num_100) và tỉ lệ bỏ qua (num_25), số bài khác nhau, số giây nghe.",
  "Người dùng không có log (22%) hoặc không có giao dịch được giữ lại với cờ has_logs và giá trị 0, vì bản thân việc thiếu hoạt động là một tín hiệu.",
  "Lấy mẫu: toàn bộ 87.330 người churn và 100.000 người ở lại, tổng 187.330. Tỉ lệ churn trong mẫu là 46,6%, KHÔNG phải tỉ lệ thật 9,0%, nên chỉ so sánh tương đối giữa các nhóm, không đọc là xác suất thật."]),
 ("4. Exploratory và Descriptive Analysis", [
  "Hình 1 và 2 cho thấy churn khác nhau rõ theo kênh đăng ký, trạng thái tự gia hạn và thời gian gắn bó.",
  "Người từng hủy (n_cancel > 0, 11% mẫu) churn 93,1% so với 40,8% ở nhóm chưa từng hủy. Người tắt tự gia hạn churn 87,8%.",
  "Kênh đăng ký 7 có churn 29,3% (70.668 người), trong khi kênh 3 và 4 có 64,6% và 72,6%. Kênh 9 là 56,2%.",
  "Có log hay không gần như không phân biệt churn (46,4% so với 47,3%), nghĩa là việc đã nghe nhạc không đủ để giữ người dùng.",
  "Kết luận cho bước mining: churn gắn chặt với hành vi đăng ký/thanh toán; hành vi nghe cho tín hiệu yếu hơn nhưng có thật."], [F / "eda_churn_by_group.png", F / "eda_churn_by_tenure.png"]),
 ("5. Data Mining Methods và Results", [
  "5.1 K-Means: chuẩn hóa sau log1p các biến hành vi; chọn k bằng silhouette trong khoảng 2–8 (k=5 cao nhất, 0,523). Ý tưởng: gán mỗi người vào tâm cụm gần nhất, lặp cho tới khi ổn định.",
  "5.2 FP-Growth: rời rạc hóa thuộc tính thành item (tự gia hạn, từng hủy, có log, mức nghe, mức bỏ qua, loại gói, kênh, giới tính), min_support = 0,02, min_lift = 1,2, chỉ giữ rule có vế phải là CHURN, rồi lọc rule dư thừa (rule dài mà rule con đã có confidence tương đương).",
  "5.3 Classification (bổ trợ): Decision Tree độ sâu 4 và Random Forest 300 cây, tập kiểm tra 25%."], [F / "clusters_pca.png"], "cluster"),
 ("6. Evaluation và Interpretation", [
  "Clustering: silhouette 0,523 ở k=5 là mức cụm tách khá rõ. Các cụm có ý nghĩa kinh doanh khác nhau rõ (bảng trên) nên diễn giải được. Điểm yếu: K-Means giả định cụm hình cầu và nhạy với outlier; cụm 3 chỉ có 1.362 người.",
  "Association rules: các rule hàng đầu có lift ≈ 2,1 (mức tối đa có thể khi churn là 46,6%) và confidence 97–99%, nhưng phần lớn chỉ phản ánh 'đã ngừng đăng ký / không có gói' tức là hệ quả của chính định nghĩa churn, không phải nguyên nhân. Dễ hiểu và dễ hành động hơn mô hình, nhưng cần đọc thận trọng.",
  "Classification: Random Forest đạt ROC-AUC 0,977 khi dùng mọi biến, chủ yếu nhờ n_cancel, auto_renew_rate, n_tx, avg_paid. Các biến này sát với định nghĩa nhãn, nên con số này đánh giá quá cao giá trị dự báo thực tế. Chỉ dùng hành vi nghe và hồ sơ, ROC-AUC còn 0,761, đó mới là mức tín hiệu từ hành vi.",
  "So sánh: clustering cho cái nhìn tổng thể và dễ trình bày; rules dễ hiểu nhưng dễ bị nhiễu do nhãn; classification chính xác nhưng kém diễn giải và dễ rò rỉ thông tin. Về quy mô, cả ba chạy tốt trên 187 nghìn người dùng; FP-Growth là bước chậm nhất khi tăng số item."]),
 ("7. Business Insights và Recommendations", [
  "Nhóm giá trị cao nên bảo vệ: cụm 'Thuê bao đang hoạt động' (65% mẫu, churn 37%, thuộc nhóm thấp trong mẫu): giữ tự gia hạn bật, nhắc gia hạn sớm.",
  "Nhóm cứu được: cụm 'Trả phí nhưng không nghe' (churn 34% trong mẫu nhưng gần như không dùng dịch vụ, 0,1 bài): kích hoạt lại bằng gợi ý nhạc cá nhân hóa, vì họ sắp thấy không đáng tiền.",
  "Nhóm 'Đã lỡ hạn nhưng vẫn nghe' (cụm 0, 11,7%, churn 95,8%): vẫn còn hoạt động nghe, nên chiến dịch kéo quay lại (ưu đãi tái đăng ký) có cơ hội tốt nhất.",
  "Kênh đăng ký 3 và 4 có churn cao nhất (65–73%): xem lại trải nghiệm onboarding và chất lượng người dùng của các kênh này; kênh 7 giữ chân tốt hơn (29%).",
  "Khuyến nghị giám sát: người dùng từng hủy và đã tắt tự gia hạn là nhóm cảnh báo sớm nhất (churn 93% và 88%).",
  "Lưu ý khi áp dụng: nhiều tín hiệu mạnh xuất hiện vì gần định nghĩa nhãn. Cần thử lại trên dữ liệu với cửa sổ thời gian tách rời giữa quan sát và dự báo để xác nhận giá trị dự báo."]),
]

# ---- docx
d = Document()
d.add_heading("DAM501 – Khám phá hành vi người dùng và mẫu rời bỏ dịch vụ nhạc đăng ký (KKBox)", 0)
d.add_paragraph("Nhóm: Lộc Nguyễn Phúc, Khánh Phạm Duy Văn, Bình Chu Đức, Quân Đoàn Đàm, Hạnh Nguyễn Thị Hồng, Tùng Trần.")
for sec in SECTIONS:
    title, paras = sec[0], sec[1]
    figs = sec[2] if len(sec) > 2 else []
    d.add_heading(title, 1)
    for p in paras:
        d.add_paragraph(p)
    if len(sec) > 3:  # cluster table
        t = d.add_table(rows=1, cols=5); t.style = "Light Grid Accent 1"
        for c, h in zip(t.rows[0].cells, ["Cụm", "Tên gọi", "Số người", "Churn (mẫu)", "Bài đã nghe"]):
            c.text = h
        for _, r in prof.iterrows():
            cells = t.add_row().cells
            cells[0].text = str(int(r.cluster)); cells[1].text = NAMES[int(r.cluster)]
            cells[2].text = f"{int(r['size']):,}"; cells[3].text = f"{r.churn_rate:.1%}"
            cells[4].text = f"{r.total_songs:.0f}"
        d.add_paragraph()
        t = d.add_table(rows=1, cols=4); t.style = "Light Grid Accent 1"
        for c, h in zip(t.rows[0].cells, ["Rule (vế trái → CHURN)", "Support", "Confidence", "Lift"]):
            c.text = h
        for _, r in rules.iterrows():
            cells = t.add_row().cells
            cells[0].text = r.antecedents; cells[1].text = f"{r.support:.3f}"
            cells[2].text = f"{r.confidence:.3f}"; cells[3].text = f"{r.lift:.2f}"
        d.add_paragraph("Feature importance (Random Forest, đầy đủ biến): " +
                        ", ".join(f"{k} {v:.2f}" for k, v in imp.items()))
    for fig in figs:
        d.add_picture(str(fig), width=Inches(5.5))
d.save(R / "report.docx")

# ---- pptx
pr = Presentation(); pr.slide_width, pr.slide_height = PI(13.33), PI(7.5)


def slide(title, bullets, img=None):
    s = pr.slides.add_slide(pr.slide_layouts[5]); s.shapes.title.text = title
    w = PI(6.3) if img else PI(12)
    tb = s.shapes.add_textbox(PI(0.6), PI(1.5), w, PI(5.5)).text_frame; tb.word_wrap = True
    for i, b in enumerate(bullets):
        p = tb.paragraphs[0] if i == 0 else tb.add_paragraph()
        p.text = "• " + b; p.font.size = Pt(18)
    if img:
        s.shapes.add_picture(str(img), PI(7.2), PI(1.6), width=PI(5.6))


slide("Khám phá hành vi và churn của dịch vụ nhạc đăng ký (KKBox)", ["DAM501 Final Project", "Dữ liệu thật: 970.960 người dùng, 9,0% churn", "Mẫu phân tích: 187.330 người dùng"])
slide("Bài toán kinh doanh", ["Doanh thu thuê bao phụ thuộc gia hạn", "Mục tiêu: hiểu nhóm hành vi và tổ hợp đặc điểm dẫn tới churn", "Kỹ thuật: K-Means, FP-Growth, Decision Tree / Random Forest"])
slide("Tín hiệu chính từ EDA", ["Từng hủy: churn 93% (so với 41%)", "Tắt tự gia hạn: churn 88%", "Kênh đăng ký 7: 29%; kênh 3 và 4: 65–73%", "Có nghe nhạc hay không: gần như không phân biệt churn"], F / "eda_churn_by_group.png")
slide("5 phân khúc người dùng (K-Means, silhouette 0,52)", [f"{NAMES[int(r.cluster)]}: {int(r['size']):,} người, churn {r.churn_rate:.0%}" for _, r in prof.iterrows()], F / "clusters_pca.png")
slide("Association rules (FP-Growth)", [f"{r.antecedents}: conf {r.confidence:.2f}, lift {r.lift:.2f}" for _, r in rules.head(5).iterrows()] + ["Lưu ý: nhiều rule phản ánh định nghĩa churn"])
slide("Classification: đọc cho đúng", ["Random Forest đủ biến: ROC-AUC 0,977 (gần nhãn)", "Chỉ hành vi nghe và hồ sơ: ROC-AUC 0,761", "Dùng để giải thích, không để khoe độ chính xác"])
slide("Khuyến nghị", ["Bảo vệ nhóm đang hoạt động: nhắc gia hạn, giữ auto-renew", "Kích hoạt lại nhóm trả phí nhưng không nghe", "Ưu đãi tái đăng ký cho nhóm đã lỡ hạn nhưng vẫn nghe", "Xem lại onboarding của kênh 3 và 4", "Hạn chế: mẫu cân bằng, nhãn gần các biến giao dịch"])
pr.save(R / "slides.pptx")
print("built")
